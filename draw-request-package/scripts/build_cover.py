#!/usr/bin/env python3
"""Generate the editable, signable Application-for-Payment cover sheet (.docx).

The cover mirrors the wording of the real Anabelle Island cover and fills in the
application number, the amount to draw, and the execution date. The land banker
sees this page first; the amount is the Exhibit D total (net of retainage).
"""
from __future__ import annotations

import argparse
import json
from datetime import date

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def ordinal(n: int) -> str:
    if 10 <= n % 100 <= 20:
        suf = "th"
    else:
        suf = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suf}"


def build(cfg: dict, total: float, out_path: str):
    app = cfg["draw_number"]
    project = cfg.get("project_name", "the Project")
    owner = cfg.get("owner", "the Owner")
    manager = cfg.get("project_manager", "Dream Finders Homes, LLC")
    agreement_date = cfg.get("agreement_date", "")
    exec_dt = cfg.get("application_date")
    d = date.fromisoformat(exec_dt) if exec_dt else date.today()
    signer = cfg.get("signer_name", "")
    signer_title = cfg.get("signer_title", "")

    doc = Document()
    for line in ("APPLICATION FOR PAYMENT", project.upper(), "CERTIFICATION FOR PAYMENT"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Application {app}").bold = True

    doc.add_paragraph(
        f'{manager} ("Project Manager"), the Project Manager under that certain '
        f'Development and Management Agreement, (the "Agreement") dated {agreement_date}, '
        f'between {owner} ("Owner"), and Project Manager, hereby requests payment '
        f"pursuant to the attached Application for Payment. Unless otherwise defined, "
        f"all capitalized terms shall have the same meaning as set forth in the "
        f"Development and Management Agreement.")

    doc.add_paragraph(
        f'Application for Payment. Attached hereto as Exhibit A is a true and correct '
        f'statement of Work performed by Contractor (the "Application for Payment"). '
        f'The Project Manager hereby requests Owner to make a payment in the amount of '
        f'${total:,.2f} (the "Payment") for the Work performed set forth in the attached '
        f'Application for Payment referred to as "Application No. {app}".',
        style="List Bullet")

    doc.add_paragraph(
        f'Certification. The undersigned individual on behalf of himself, individually, '
        f'and as an authorized agent or officer of "Project Manager" hereby certifies '
        f'that the foregoing information is true and correct of his own knowledge. '
        f'Executed as of {ordinal(d.day)} day of {d.strftime("%B")}, {d.year}.',
        style="List Bullet")

    doc.add_paragraph("\n\n")
    doc.add_paragraph("By: ____________________________")
    doc.add_paragraph(f"Name: {signer}")
    doc.add_paragraph(f"Title: {signer_title}")
    doc.add_paragraph(f"Date: {d.strftime('%-m/%-d/%Y')}")
    doc.save(out_path)
    print(f"Cover sheet -> {out_path}  (Application {app}, ${total:,.2f})")


def main():
    ap = argparse.ArgumentParser(description="Generate the Application-for-Payment cover sheet .docx")
    ap.add_argument("invoices_json", help="invoices.json (for draw #, project, dates, total)")
    ap.add_argument("out", help="Output .docx path")
    ap.add_argument("--total", type=float, help="Override draw total (default: sum of net amounts)")
    args = ap.parse_args()
    with open(args.invoices_json) as fh:
        cfg = json.load(fh)
    total = args.total
    if total is None:
        total = round(sum(i.get("net_amount", i["gross_amount"] * (1 - i.get("retainage_rate", 0)))
                          for i in cfg["invoices"]), 2)
    build(cfg, total, args.out)


if __name__ == "__main__":
    main()
