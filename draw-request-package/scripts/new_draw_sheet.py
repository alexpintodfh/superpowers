#!/usr/bin/env python3
"""Create the next 'Draw Request N+1' sheet from the latest draw sheet.

Mirrors how the tracker grows: the finished draw's Current Period column freezes
into a new 'Draw N' history column, every summary column shifts one to the right,
and the new Current Period / Retainage columns reset to zero — ready for the next
batch of invoices. Summary formulas are rewritten fresh (openpyxl does not adjust
formula references across an inserted column).
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as col_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import draw_lib as dl


def _shift_merges(ws, insert_col):
    """Shift/extend merged ranges to match a column inserted at insert_col.

    openpyxl.insert_cols moves cell *values* but leaves merged ranges in place, so
    a value shifted out of its merge anchor is dropped on save. Rebuild the ranges:
    columns fully at/after the insert shift right by one; a range straddling the
    insert grows by one.
    """
    from openpyxl.worksheet.cell_range import CellRange
    old = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    ws.merged_cells.ranges = []
    for cr in old:
        min_c, max_c = cr.min_col, cr.max_col
        if min_c >= insert_col:
            min_c += 1
            max_c += 1
        elif max_c >= insert_col:
            max_c += 1
        ws.merge_cells(start_row=cr.min_row, start_column=min_c,
                       end_row=cr.max_row, end_column=max_c)


def _update_title_date(ws, new_number):
    """Refresh a 'Draw Request #N - M/D/YYYY' style title cell in the top rows.

    The signature date is always the 10th of the current month.
    """
    sig = date.today().replace(day=10)
    slash = f"{sig.month}/{sig.day}/{sig.year}"
    for r in range(1, 4):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "draw" in v.lower() and "#" in v:
                v2 = re.sub(r"#\s*\d+", f"#{new_number}", v)
                v2 = re.sub(r"\d{1,2}/\d{1,2}/\d{4}", slash, v2)
                ws.cell(row=r, column=c).value = v2
                return


def _refresh_row_formulas(ws, cols):
    for r in range(dl.FIRST_COST_ROW, dl.LAST_COST_ROW + 1):
        for role, formula in dl.summary_formulas(r, cols).items():
            ws.cell(row=r, column=cols[role]).value = formula


def _refresh_total_row(ws, cols):
    """Rewrite the TOTAL COSTS row sums for the money columns after the shift."""
    # find the TOTAL COSTS row
    total_row = None
    for r in range(dl.LAST_COST_ROW + 1, ws.max_row + 1):
        if dl._norm(ws.cell(row=r, column=2).value) == "total costs":
            total_row = r
            break
    if total_row is None:
        return
    money_cols = [dl.CURRENT_BUDGET_COL]                  # F (current budget)
    money_cols += list(range(dl.FIRST_DRAW_COL, dl.last_draw_col(cols) + 1))  # all draw history
    money_cols += [cols["prev"], cols["cp"], cols["cplr"], cols["tcd"], cols["rtd"], cols["bal"]]
    for c in money_cols:
        L = col_letter(c)
        ws.cell(row=total_row, column=c).value = f"=SUM({L}5:{L}50)"


def _reset_embedded_exhibit(ws, new_number):
    head = None
    for r in range(1, ws.max_row + 1):
        if dl._norm(ws.cell(row=r, column=2).value) == 'exhibit "d"':
            head = r
            break
    if head is None:
        return
    ws.cell(row=head + 1, column=6, value=new_number)  # Draw Request # value (col F)
    for r in range(head + 4, ws.max_row + 2):
        first = ws.cell(row=r, column=2).value
        for c in range(2, 7):
            ws.cell(row=r, column=c).value = None
        if dl._norm(first) == "total":
            break


def create_next_draw(workbook, out=None, new_number=None):
    wb = openpyxl.load_workbook(workbook)
    src, n = dl.latest_draw_sheet(wb)
    if new_number is None:
        new_number = n + 1
    naming = dl.draw_sheet_naming(wb)            # match this workbook's style ('Draw' vs 'Draw Request')
    new_title = f"{naming} {new_number}"
    if any(dl.draw_sheet_number(s) == new_number for s in wb.sheetnames):
        sys.exit(f"A draw sheet for #{new_number} already exists.")

    ws = wb.copy_worksheet(src)
    ws.title = new_title

    cols = dl.find_columns(ws)
    prev_col = cols["prev"]

    # Insert the frozen-history column just left of TOTAL PREVIOUSLY DRAWN.
    ws.insert_cols(prev_col, 1)
    _shift_merges(ws, prev_col)              # openpyxl moves values but not merges
    cols = dl.find_columns(ws)               # re-locate after the shift
    history_col = dl.last_draw_col(cols)     # == the inserted column

    # Freeze the finished draw's Current Period (gross) into the new history column.
    ws.cell(row=dl.HEADER_ROW, column=history_col, value=f"Draw {n}")
    src_fmt = ws.cell(row=dl.FIRST_COST_ROW, column=dl.FIRST_DRAW_COL).number_format
    for r in range(dl.FIRST_COST_ROW, dl.LAST_COST_ROW + 1):
        frozen = ws.cell(row=r, column=cols["cp"]).value
        ws.cell(row=r, column=history_col, value=frozen if isinstance(frozen, (int, float)) else 0)
        ws.cell(row=r, column=history_col).number_format = src_fmt
        # reset the new input columns
        ws.cell(row=r, column=cols["cp"]).value = 0
        ws.cell(row=r, column=cols["ret"]).value = 0

    _refresh_row_formulas(ws, cols)
    _refresh_total_row(ws, cols)
    _reset_embedded_exhibit(ws, new_number)
    _update_title_date(ws, new_number)

    # point the standalone Exhibit D at the new draw and clear its lines
    ex_name = next((s for s in wb.sheetnames if dl._norm(s) == "exhibit d"), None)
    if ex_name:
        ex = wb[ex_name]
        ex["E2"] = new_number
        for r in range(5, ex.max_row + 2):
            for c in range(1, 6):
                ex.cell(row=r, column=c).value = None

    out = out or workbook
    wb.save(out)
    print(f"Created '{new_title}' (froze Draw {n} into history, reset Current Period).")
    print(f"Saved -> {out}")
    return new_number


def main():
    ap = argparse.ArgumentParser(description="Add the next draw sheet to the tracker.")
    ap.add_argument("workbook")
    ap.add_argument("--out")
    ap.add_argument("--number", type=int, help="Draw number to create (default: latest + 1)")
    args = ap.parse_args()
    create_next_draw(args.workbook, args.out, args.number)


if __name__ == "__main__":
    main()
