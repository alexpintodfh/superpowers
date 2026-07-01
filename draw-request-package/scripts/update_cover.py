#!/usr/bin/env python3
"""Update a community's PRIOR cover sheet .docx into the new draw's cover.

Cover formats differ per community/land banker, so the reliable approach is to
take the previous draw's cover as a template and substitute three things:
  * the draw/application number (#N),
  * the signature date (always the 10th of the current month), and
  * the draw amount (the Exhibit D net total).

Only paragraphs that actually contain one of these is rewritten, so the rest of
the community-specific wording and formatting is preserved.
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import date

from docx import Document

MONTHS = ("January February March April May June July August September "
          "October November December").split()
DATE_RE = re.compile(
    r"(" + "|".join(MONTHS) + r")\s+(\d{1,2})(st|nd|rd|th)?,\s*(\d{4})")
SLASH_DATE_RE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")
CURRENCY_RE = re.compile(r"\$\s?[\d,]+\.\d{2}")


def _ordinal(n: int) -> str:
    if 10 <= n % 100 <= 20:
        return f"{n}th"
    return f"{n}" + {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")


def signature_date(today=None) -> date:
    today = today or date.today()
    return today.replace(day=10)


def _para_text(p) -> str:
    return "".join(r.text for r in p.runs)


def _set_para_text(p, text, bold_substr=None):
    """Replace a paragraph's text, keeping the first run's font.

    If bold_substr is given, that substring is emitted as its own bold run so the
    draw amount shows bold in the cover (matches the rest of the document's style).
    """
    base = p.runs[0] if p.runs else None
    base_name = base.font.name if base else None
    base_size = base.font.size if base else None

    def _new_run(seg, bold):
        run = p.add_run(seg)
        if base_name:
            run.font.name = base_name
        if base_size:
            run.font.size = base_size
        if bold:
            run.bold = True
        return run

    for r in list(p.runs):
        r.text = ""
    # remove the now-empty original runs so only our rebuilt runs remain
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

    if bold_substr and bold_substr in text:
        before, after = text.split(bold_substr, 1)
        if before:
            _new_run(before, False)
        _new_run(bold_substr, True)
        if after:
            _new_run(after, False)
    else:
        _new_run(text, False)


def transform_text(text, old_number, new_number, sig, total):
    long_date = f"{sig.strftime('%B')} {sig.day}, {sig.year}"
    long_date_ord = f"{sig.strftime('%B')} {_ordinal(sig.day)}, {sig.year}"
    slash_date = f"{sig.month}/{sig.day}/{sig.year}"

    # draw / application number: "#13" -> "#14"
    text = re.sub(rf"#\s*{old_number}\b", f"#{new_number}", text)

    # Dates: only the signature/draw date changes. Leave fixed reference dates
    # (e.g. the Construction/Development Agreement date) untouched — those live in
    # a sentence mentioning the agreement, so skip date edits in such paragraphs.
    if "agreement" not in text.lower():
        def _date_sub(m):
            return long_date_ord if m.group(3) else long_date
        text = DATE_RE.sub(_date_sub, text)
        text = SLASH_DATE_RE.sub(slash_date, text)

    # currency: only meaningful in an "amount" sentence
    if "amount" in text.lower():
        text = CURRENCY_RE.sub(f"${total:,.2f}", text)
    return text


def update_cover(template, invoices_json, out_path, old_number=None, total=None, today=None):
    with open(invoices_json) as fh:
        cfg = json.load(fh)
    new_number = cfg["draw_number"]
    if old_number is None:
        old_number = new_number - 1
    if total is None:
        total = round(sum(i.get("net_amount", i["gross_amount"] * (1 - i.get("retainage_rate", 0)))
                          for i in cfg["invoices"]), 2)
    sig = signature_date(today)
    amount_str = f"${total:,.2f}"

    def _apply(p):
        original = _para_text(p)
        updated = transform_text(original, old_number, new_number, sig, total)
        if updated != original:
            # bold the draw amount when it appears in this paragraph
            bold = amount_str if amount_str in updated else None
            _set_para_text(p, updated, bold_substr=bold)

    doc = Document(template)
    for p in doc.paragraphs:
        _apply(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _apply(p)
    doc.save(out_path)
    print(f"Cover sheet -> {out_path}  (#{new_number}, {sig.strftime('%-m/%-d/%Y')}, ${total:,.2f})")


def main():
    ap = argparse.ArgumentParser(description="Update a prior cover .docx into the new draw's cover.")
    ap.add_argument("template", help="Prior draw's cover .docx")
    ap.add_argument("invoices_json")
    ap.add_argument("out", help="Output .docx path")
    ap.add_argument("--old-number", type=int, help="Draw number in the template (default: new-1)")
    ap.add_argument("--total", type=float, help="Override draw total")
    args = ap.parse_args()
    update_cover(args.template, args.invoices_json, args.out, args.old_number, args.total)


if __name__ == "__main__":
    main()
